#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import os
import subprocess
import sys
from pathlib import Path
from typing import Any


def script_dir() -> Path:
    return Path(__file__).resolve().parent


def venv_python(repo_root: Path) -> Path:
    if os.name == "nt":
        return repo_root / ".armarius" / "venv" / "Scripts" / "python.exe"
    return repo_root / ".armarius" / "venv" / "bin" / "python"


def maybe_reexec_in_venv(args: argparse.Namespace) -> None:
    if args.no_bootstrap or os.environ.get("ARMARIUS_IN_VENV") == "1":
        return
    repo_root = Path(args.repo_root or os.getcwd()).resolve()
    completed = subprocess.run(
        [
            sys.executable,
            str(script_dir() / "armarius_bootstrap.py"),
            "--repo-root",
            str(repo_root),
            "--profile",
            "write",
            "--json",
        ],
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
    )
    if completed.returncode != 0:
        sys.stderr.write(completed.stdout)
        sys.stderr.write(completed.stderr)
        raise subprocess.CalledProcessError(completed.returncode, completed.args)
    python_path = venv_python(repo_root)
    if python_path.exists():
        env = dict(os.environ)
        env["ARMARIUS_IN_VENV"] = "1"
        os.execve(str(python_path), [str(python_path), __file__, *sys.argv[1:]], env)


def default_output(repo_root: Path, filename: str) -> Path:
    out_dir = repo_root / ".armarius" / "outputs"
    out_dir.mkdir(parents=True, exist_ok=True)
    return out_dir / filename


def read_input_text(path: str | None) -> str:
    if path:
        return Path(path).read_text(encoding="utf-8")
    return sys.stdin.read()


def write_docx(text: str, output: Path) -> None:
    import docx

    document = docx.Document()
    for block in text.splitlines():
        if block.startswith("# "):
            document.add_heading(block[2:].strip(), level=1)
        elif block.startswith("## "):
            document.add_heading(block[3:].strip(), level=2)
        elif block.strip():
            document.add_paragraph(block)
        else:
            document.add_paragraph("")
    document.save(str(output))


def write_pdf(text: str, output: Path) -> None:
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(output), pagesize=letter)
    width, height = letter
    x = 72
    y = height - 72
    line_height = 14
    for raw_line in text.splitlines() or [""]:
        line = raw_line[:110]
        if y < 72:
            c.showPage()
            y = height - 72
        c.drawString(x, y, line)
        y -= line_height
    c.save()


def rows_from_json(path: Path) -> list[list[Any]]:
    data = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(data, list):
        if not data:
            return []
        if all(isinstance(item, dict) for item in data):
            keys = sorted({key for item in data for key in item})
            return [keys] + [[item.get(key) for key in keys] for item in data]
        if all(isinstance(item, list) for item in data):
            return data
    if isinstance(data, dict) and isinstance(data.get("content"), list):
        rows = [["type", "locator", "text"]]
        for item in data["content"]:
            if isinstance(item, dict):
                rows.append([item.get("type"), item.get("locator"), item.get("text")])
        return rows
    raise ValueError("JSON input must be a list of rows, list of objects, or Armarius extraction JSON.")


def write_xlsx(rows: list[list[Any]], output: Path) -> None:
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    for row in rows:
        sheet.append(row)
    workbook.save(str(output))


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Create basic documents from text or JSON.")
    parser.add_argument("--repo-root", help="Target repository root. Defaults to the current directory.")
    parser.add_argument("--no-bootstrap", action="store_true", help="Do not create/use .armarius/venv first.")
    subparsers = parser.add_subparsers(dest="command", required=True)

    docx_parser = subparsers.add_parser("text-to-docx", help="Write plain/markdown-like text to DOCX.")
    docx_parser.add_argument("--input", help="Input text path. Defaults to stdin.")
    docx_parser.add_argument("--output", help="Output DOCX path. Defaults to .armarius/outputs/output.docx.")

    pdf_parser = subparsers.add_parser("text-to-pdf", help="Write plain text to a simple PDF.")
    pdf_parser.add_argument("--input", help="Input text path. Defaults to stdin.")
    pdf_parser.add_argument("--output", help="Output PDF path. Defaults to .armarius/outputs/output.pdf.")

    xlsx_parser = subparsers.add_parser("json-to-xlsx", help="Write JSON rows or extraction content to XLSX.")
    xlsx_parser.add_argument("--input", required=True, help="Input JSON path.")
    xlsx_parser.add_argument("--output", help="Output XLSX path. Defaults to .armarius/outputs/output.xlsx.")
    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()
    maybe_reexec_in_venv(args)
    repo_root = Path(args.repo_root or os.getcwd()).resolve()

    if args.command == "text-to-docx":
        output = Path(args.output).resolve() if args.output else default_output(repo_root, "output.docx")
        output.parent.mkdir(parents=True, exist_ok=True)
        write_docx(read_input_text(args.input), output)
    elif args.command == "text-to-pdf":
        output = Path(args.output).resolve() if args.output else default_output(repo_root, "output.pdf")
        output.parent.mkdir(parents=True, exist_ok=True)
        write_pdf(read_input_text(args.input), output)
    elif args.command == "json-to-xlsx":
        output = Path(args.output).resolve() if args.output else default_output(repo_root, "output.xlsx")
        output.parent.mkdir(parents=True, exist_ok=True)
        write_xlsx(rows_from_json(Path(args.input)), output)
    else:
        parser.error(f"unknown command: {args.command}")

    print(f"output={output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
