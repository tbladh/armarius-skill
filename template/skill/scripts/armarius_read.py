#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import importlib.util
import json
import os
import re
import subprocess
import sys
from pathlib import Path
from typing import Any


TEXT_EXTENSIONS = {".txt", ".md", ".rst", ".log", ".json", ".xml", ".yaml", ".yml"}
POWERPOINT_EXTENSIONS = {".pptx", ".pptm", ".ppsx", ".ppsm", ".potx", ".potm"}


def module_available(name: str) -> bool:
    return importlib.util.find_spec(name) is not None


def script_dir() -> Path:
    return Path(__file__).resolve().parent


def venv_python(repo_root: Path) -> Path:
    if os.name == "nt":
        return repo_root / ".armarius" / "venv" / "Scripts" / "python.exe"
    return repo_root / ".armarius" / "venv" / "bin" / "python"


def maybe_reexec_in_venv(args: argparse.Namespace) -> None:
    if args.no_bootstrap or os.environ.get("ARMARIUS_IN_VENV") == "1":
        return
    bootstrap = script_dir() / "armarius_bootstrap.py"
    repo_root = Path(args.repo_root or os.getcwd()).resolve()
    completed = subprocess.run(
        [sys.executable, str(bootstrap), "--repo-root", str(repo_root), "--profile", args.profile, "--json"],
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
    )
    if completed.returncode != 0:
        sys.stderr.write(completed.stdout)
        sys.stderr.write(completed.stderr)
        raise subprocess.CalledProcessError(completed.returncode, completed.args)
    python_path = venv_python(repo_root)
    if not python_path.exists():
        return
    env = dict(os.environ)
    env["ARMARIUS_IN_VENV"] = "1"
    os.execve(str(python_path), [str(python_path), __file__, *sys.argv[1:]], env)


def stable_output_dir(repo_root: Path, source: Path, explicit: str | None) -> Path:
    if explicit:
        return Path(explicit).resolve()
    stem = re.sub(r"[^A-Za-z0-9._-]+", "-", source.stem).strip("-") or "document"
    return repo_root / ".armarius" / "outputs" / stem


def base_result(source: Path, kind: str) -> dict[str, Any]:
    return {
        "source_path": str(source),
        "kind": kind,
        "metadata": {},
        "content": [],
        "warnings": [],
        "missing_dependencies": [],
        "outputs": {},
    }


def add_missing(result: dict[str, Any], package: str, purpose: str) -> None:
    result["missing_dependencies"].append({"package": package, "purpose": purpose})


def read_pdf(source: Path) -> dict[str, Any]:
    result = base_result(source, "pdf")
    if module_available("pypdf"):
        from pypdf import PdfReader

        reader = PdfReader(str(source))
        result["metadata"]["page_count"] = len(reader.pages)
        if reader.metadata:
            result["metadata"]["pdf_metadata"] = {
                str(key): str(value) for key, value in dict(reader.metadata).items()
            }
        for index, page in enumerate(reader.pages, start=1):
            try:
                text = page.extract_text() or ""
            except Exception as exc:
                text = ""
                result["warnings"].append(f"pypdf failed on page {index}: {exc}")
            result["content"].append({"type": "page", "locator": f"page {index}", "text": text})
    else:
        add_missing(result, "pypdf", "basic PDF text extraction")

    if module_available("pdfplumber"):
        import pdfplumber

        try:
            with pdfplumber.open(str(source)) as pdf:
                for page_index, page in enumerate(pdf.pages, start=1):
                    for table_index, table in enumerate(page.extract_tables() or [], start=1):
                        result["content"].append(
                            {
                                "type": "table",
                                "locator": f"page {page_index} table {table_index}",
                                "rows": table,
                            }
                        )
        except Exception as exc:
            result["warnings"].append(f"pdfplumber table extraction failed: {exc}")
    else:
        add_missing(result, "pdfplumber", "PDF layout and table extraction")

    text_len = sum(len(item.get("text", "")) for item in result["content"])
    page_count = int(result["metadata"].get("page_count") or 0)
    if page_count and text_len < page_count * 40:
        result["warnings"].append("Low extracted text volume; the PDF may be scanned or image-heavy.")
    return result


def read_docx(source: Path) -> dict[str, Any]:
    result = base_result(source, "docx")
    if not module_available("docx"):
        add_missing(result, "python-docx", "DOCX paragraph and table extraction")
        return result

    import docx

    doc = docx.Document(str(source))
    result["metadata"]["paragraph_count"] = len(doc.paragraphs)
    result["metadata"]["table_count"] = len(doc.tables)
    for index, paragraph in enumerate(doc.paragraphs, start=1):
        text = paragraph.text or ""
        if text:
            result["content"].append(
                {"type": "paragraph", "locator": f"paragraph {index}", "text": text}
            )
    for table_index, table in enumerate(doc.tables, start=1):
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        result["content"].append({"type": "table", "locator": f"table {table_index}", "rows": rows})
    return result


def read_pptx(source: Path) -> dict[str, Any]:
    result = base_result(source, "pptx")
    if not module_available("pptx"):
        add_missing(result, "python-pptx", "PPTX slide, text, and table extraction")
        return result

    from pptx import Presentation

    presentation = Presentation(str(source))
    result["metadata"]["slide_count"] = len(presentation.slides)
    result["metadata"]["slide_width"] = presentation.slide_width
    result["metadata"]["slide_height"] = presentation.slide_height

    for slide_index, slide in enumerate(presentation.slides, start=1):
        slide_text: list[str] = []
        title = ""
        table_index = 0

        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame:
                text = shape.text_frame.text or ""
                if text:
                    slide_text.append(text)
                    if shape == slide.shapes.title:
                        title = text
            if getattr(shape, "has_table", False):
                table_index += 1
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                result["content"].append(
                    {
                        "type": "table",
                        "locator": f"slide {slide_index} table {table_index}",
                        "rows": rows,
                    }
                )

        notes_text = ""
        try:
            if getattr(slide, "has_notes_slide", False):
                notes_text = slide.notes_slide.notes_text_frame.text or ""
        except Exception as exc:
            result["warnings"].append(f"speaker notes extraction failed on slide {slide_index}: {exc}")

        result["content"].append(
            {
                "type": "slide",
                "locator": f"slide {slide_index}",
                "title": title,
                "text": "\n".join(slide_text),
                "notes": notes_text,
            }
        )
    return result


def read_workbook(source: Path) -> dict[str, Any]:
    ext = source.suffix.lower()
    result = base_result(source, "xlsx" if ext in {".xlsx", ".xlsm"} else "xls")
    if ext == ".xls" and not module_available("xlrd"):
        add_missing(result, "xlrd", "legacy XLS extraction")
        result["warnings"].append("LibreOffice conversion to XLSX may produce better legacy XLS results.")
        return result
    if ext in {".xlsx", ".xlsm"} and not module_available("openpyxl"):
        add_missing(result, "openpyxl", "XLSX extraction")
        return result

    if ext == ".xls":
        import xlrd

        book = xlrd.open_workbook(str(source))
        result["metadata"]["sheet_names"] = book.sheet_names()
        for sheet in book.sheets():
            rows = [sheet.row_values(i) for i in range(sheet.nrows)]
            result["content"].append(
                {
                    "type": "sheet",
                    "locator": sheet.name,
                    "rows": rows,
                    "row_count": sheet.nrows,
                    "column_count": sheet.ncols,
                }
            )
        return result

    import openpyxl

    wb_values = openpyxl.load_workbook(str(source), data_only=True, read_only=True)
    wb_formulas = openpyxl.load_workbook(str(source), data_only=False, read_only=True)
    result["metadata"]["sheet_names"] = wb_values.sheetnames
    for name in wb_values.sheetnames:
        sheet_values = wb_values[name]
        sheet_formulas = wb_formulas[name]
        rows = []
        formulas = []
        for row_index, row in enumerate(sheet_values.iter_rows(values_only=True), start=1):
            rows.append(list(row))
            formula_row = []
            for cell in sheet_formulas[row_index]:
                value = cell.value
                formula_row.append(value if isinstance(value, str) and value.startswith("=") else None)
            if any(value is not None for value in formula_row):
                formulas.append({"row": row_index, "formulas": formula_row})
        result["content"].append(
            {
                "type": "sheet",
                "locator": name,
                "rows": rows,
                "formulas": formulas,
                "row_count": sheet_values.max_row,
                "column_count": sheet_values.max_column,
            }
        )
    return result


def read_delimited(source: Path, delimiter: str) -> dict[str, Any]:
    result = base_result(source, "tsv" if delimiter == "\t" else "csv")
    with source.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.reader(handle, delimiter=delimiter))
    result["content"].append(
        {
            "type": "table",
            "locator": source.name,
            "rows": rows,
            "row_count": len(rows),
            "column_count": max((len(row) for row in rows), default=0),
        }
    )
    return result


def read_html(source: Path) -> dict[str, Any]:
    result = base_result(source, "html")
    text = source.read_text(encoding="utf-8", errors="replace")
    if module_available("bs4"):
        from bs4 import BeautifulSoup

        soup = BeautifulSoup(text, "html.parser")
        result["metadata"]["title"] = soup.title.string if soup.title and soup.title.string else None
        result["content"].append({"type": "text", "locator": source.name, "text": soup.get_text("\n")})
    else:
        add_missing(result, "beautifulsoup4", "HTML text extraction")
        result["content"].append({"type": "text", "locator": source.name, "text": text})
    return result


def read_text(source: Path) -> dict[str, Any]:
    result = base_result(source, "text")
    result["content"].append(
        {"type": "text", "locator": source.name, "text": source.read_text(encoding="utf-8", errors="replace")}
    )
    return result


def read_document(source: Path) -> dict[str, Any]:
    ext = source.suffix.lower()
    if ext == ".pdf":
        return read_pdf(source)
    if ext == ".docx":
        return read_docx(source)
    if ext == ".doc":
        result = base_result(source, "doc")
        result["warnings"].append("Legacy DOC needs conversion, usually with LibreOffice/soffice, before Python extraction.")
        return result
    if ext in POWERPOINT_EXTENSIONS:
        return read_pptx(source)
    if ext in {".ppt", ".pps", ".pot"}:
        result = base_result(source, "ppt")
        result["warnings"].append("Legacy PPT needs conversion, usually with LibreOffice/soffice, before Python extraction.")
        return result
    if ext in {".xlsx", ".xlsm", ".xls"}:
        return read_workbook(source)
    if ext == ".csv":
        return read_delimited(source, ",")
    if ext == ".tsv":
        return read_delimited(source, "\t")
    if ext in {".html", ".htm"}:
        return read_html(source)
    if ext in TEXT_EXTENSIONS:
        return read_text(source)
    result = base_result(source, "unknown")
    result["warnings"].append(f"Unknown extension {ext or '<none>'}; attempted UTF-8 text read.")
    try:
        result["content"].append({"type": "text", "locator": source.name, "text": source.read_text(encoding="utf-8", errors="replace")})
    except Exception as exc:
        result["warnings"].append(f"Could not read as text: {exc}")
    return result


def text_from_result(result: dict[str, Any]) -> str:
    lines = [
        f"Source: {result['source_path']}",
        f"Kind: {result['kind']}",
        "",
    ]
    for warning in result["warnings"]:
        lines.append(f"WARNING: {warning}")
    if result["warnings"]:
        lines.append("")
    for item in result["content"]:
        locator = item.get("locator", "content")
        item_type = item.get("type", "item")
        lines.append(f"## {locator} ({item_type})")
        if "text" in item:
            lines.append(str(item["text"]))
        if "notes" in item and item["notes"]:
            lines.append("")
            lines.append("Notes:")
            lines.append(str(item["notes"]))
        if "rows" in item:
            for row in item["rows"]:
                lines.append("\t".join("" if value is None else str(value) for value in row))
        lines.append("")
    return "\n".join(lines)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Extract a document into Armarius JSON and text artifacts.")
    parser.add_argument("source", help="Document path to extract.")
    parser.add_argument("--repo-root", help="Target repository root. Defaults to the current directory.")
    parser.add_argument("--out", help="Output directory. Defaults to .armarius/outputs/<document-stem>.")
    parser.add_argument("--profile", default="read", choices=["read", "all", "ocr"], help="Bootstrap profile.")
    parser.add_argument("--no-bootstrap", action="store_true", help="Do not create/use .armarius/venv first.")
    parser.add_argument("--json", action="store_true", help="Print JSON result.")
    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()
    maybe_reexec_in_venv(args)

    source = Path(args.source).resolve()
    if not source.exists():
        parser.error(f"source does not exist: {source}")
    repo = Path(args.repo_root or os.getcwd()).resolve()
    out_dir = stable_output_dir(repo, source, args.out)
    out_dir.mkdir(parents=True, exist_ok=True)

    result = read_document(source)
    result["outputs"] = {
        "json": str(out_dir / "extraction.json"),
        "text": str(out_dir / "extracted.txt"),
    }
    text = text_from_result(result)
    (out_dir / "extraction.json").write_text(json.dumps(result, indent=2, sort_keys=True), encoding="utf-8")
    (out_dir / "extracted.txt").write_text(text, encoding="utf-8")

    if args.json:
        print(json.dumps(result, indent=2, sort_keys=True))
    else:
        print(f"json={out_dir / 'extraction.json'}")
        print(f"text={out_dir / 'extracted.txt'}")
        if result["warnings"]:
            print(f"warnings={len(result['warnings'])}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
